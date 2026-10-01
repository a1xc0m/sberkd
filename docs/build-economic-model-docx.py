from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "economic-model-calculation-2026.docx"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    bd = tc_pr.find(qn("w:tcBorders"))
    if bd is None:
        bd = OxmlElement("w:tcBorders")
        tc_pr.append(bd)
    for edge in ("top", "left", "bottom", "right"):
        el = bd.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            bd.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        child = mar.find(qn("w:" + side))
        if child is None:
            child = OxmlElement("w:" + side)
            mar.append(child)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, right_align_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for j, header in enumerate(headers):
        table.columns[j].width = Inches(widths[j])
        cell = table.rows[0].cells[j]
        cell.width = Inches(widths[j])
        cell.text = header
        shade(cell, "253D3A")
        borders(cell)
        cell_margin(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if j in right_align_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            cell.width = Inches(widths[j])
            cell.text = str(value)
            shade(cell, "F4F7F6" if i % 2 else "FFFFFF")
            borders(cell)
            cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                if j in right_align_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(25, 34, 32)
                    if i == len(rows) - 1 and value in ("−9,51 млн ₽", "+28,99 млн ₽"):
                        run.font.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Cm(1.75)
sec.bottom_margin = Cm(1.65)
sec.left_margin = Cm(2.05)
sec.right_margin = Cm(2.05)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(25, 34, 32)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.14

title_style = doc.styles["Title"]
title_style.font.name = "Arial"
title_style.font.size = Pt(19)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(7)
for border in title_style._element.get_or_add_pPr().findall(qn("w:pBdr")):
    title_style._element.get_or_add_pPr().remove(border)

h1 = doc.styles["Heading 1"]
h1.font.name = "Arial"
h1.font.size = Pt(13)
h1.font.bold = True
h1.font.color.rgb = RGBColor(0, 0, 0)
h1.paragraph_format.space_before = Pt(14)
h1.paragraph_format.space_after = Pt(6)
h1.paragraph_format.keep_with_next = True

title = doc.add_paragraph("Расчёт экономической модели", "Title")
for border in title._p.get_or_add_pPr().findall(qn("w:pBdr")):
    title._p.get_or_add_pPr().remove(border)
para(doc, "Семейный конструктор целей. Проверка модели презентации от 30 сентября 2026 года.")
para(doc, "Модель показывает возможный вклад в экономику экосистемы при заданных параметрах. Из публичных данных подтверждён MAU СберKids; конверсия, чек, доля дополнительных продаж, маржа и затраты пока требуют проверки. Значение +29 млн ₽ во второй год нельзя представлять как подтверждённый доход Банка.")

doc.add_paragraph("Исходные параметры", "Heading 1")
add_table(
    doc,
    ["Параметр", "Значение", "Основание"],
    [
        ("MAU СберKids, май 2026", "2,9 млн", "Публичное раскрытие Сбера [1]"),
        ("Активных детей на семью", "1,25", "Допущение для устранения повторного счёта семей"),
        ("Подключение к концу года 1", "15%", "Гипотеза активации продукта"),
        ("Средняя база в году 1", "50%", "Линейный набор аудитории от нуля"),
        ("Семьи с оплатой подарка за год", "45%", "Гипотеза полной воронки до покупки"),
        ("Средний чек подарка", "5 500 ₽", "Гипотеза по подарочной корзине"),
        ("Дополнительный оборот", "60%", "Гипотеза причинного эффекта"),
        ("Вклад после переменных затрат", "8%", "Гипотеза маржи от дополнительного оборота"),
        ("Разработка / годовое содержание", "18 / 12 млн ₽", "Предварительные бюджетные заглушки"),
    ],
    [2.65, 1.15, 3.1],
)

doc.add_paragraph("Логика расчёта", "Heading 1")
para(doc, "Потенциальная база: 2,9 млн MAU / 1,25 ребёнка на семью = 2,32 млн. В презентации для простоты использованы 2,3 млн семей; все дальнейшие числа рассчитаны от этой округлённой базы. Если считать без промежуточного округления, подключённых семей получится 348 тыс. вместо 345 тыс.")
para(doc, "Подключённые семьи к концу первого года: 2,3 млн × 15% = 345 тыс. При линейном росте средняя активная база первого года — 172,5 тыс. семей. Во втором году принята постоянная база 345 тыс. без роста и оттока.")
para(doc, "Формула результата: средняя активная база × 45% покупающих семей × 5 500 ₽ × 60% дополнительного оборота × 8% маржинального вклада − затраты периода.")

doc.add_page_break()
doc.add_paragraph("Расчёт по годам", "Heading 1")
add_table(
    doc,
    ["Показатель", "Год 1", "Год 2"],
    [
        ("Среднее число активных семей", "172 500", "345 000"),
        ("Оплаченные подарки", "77 625", "155 250"),
        ("Оборот подарков", "426,94 млн ₽", "853,88 млн ₽"),
        ("Дополнительный оборот", "256,16 млн ₽", "512,33 млн ₽"),
        ("Вклад после переменных затрат", "20,49 млн ₽", "40,99 млн ₽"),
        ("Разработка и содержание", "30,00 млн ₽", "12,00 млн ₽"),
        ("Результат периода", "−9,51 млн ₽", "+28,99 млн ₽"),
    ],
    [2.9, 2.0, 2.0],
    right_align_cols=(1, 2),
)
para(doc, "Совокупный результат двух лет — +19,48 млн ₽. В первый год учтены 18 млн ₽ разработки и 12 млн ₽ содержания; во второй — 12 млн ₽ содержания. Это расчёт для экосистемы, а не распределение дохода между Мегамаркетом и Банком.")

doc.add_paragraph("Порог и чувствительность", "Heading 1")
para(doc, "Для покрытия 12 млн ₽ годовых затрат во втором году при чеке 5 500 ₽ необходимо 12 000 000 / (5 500 × 60% × 8%) ≈ 45 455 оплаченных подарков. Для возмещения всех 42 млн ₽ затрат за два года при остальных базовых параметрах нужен средний чек около 3 760 ₽.")
add_table(
    doc,
    ["Средний чек", "Год 1", "Год 2", "Два года"],
    [
        ("1 500 ₽", "−24,4 млн ₽", "−0,8 млн ₽", "−25,2 млн ₽"),
        ("3 000 ₽", "−18,8 млн ₽", "+10,4 млн ₽", "−8,5 млн ₽"),
        ("5 500 ₽", "−9,5 млн ₽", "+29,0 млн ₽", "+19,5 млн ₽"),
    ],
    [1.25, 1.85, 1.85, 1.95],
    right_align_cols=(1, 2, 3),
)
para(doc, "При чеке 3 000 ₽ второй год положителен, но двухлетние вложения ещё не возвращаются. Остальные параметры в таблице неизменны; это чувствительность к одному фактору, а не новый прогноз.")

doc.add_page_break()
doc.add_paragraph("Обоснование и ограничения", "Heading 1")
para(doc, "Публичный факт — 2,9 млн MAU СберKids в мае 2026 года [1]. Это число детей, пользовавшихся приложением за месяц, а не число семей с подтверждённой связью со взрослым. Данные по подросткам 14–17 лет в формулу не включены. Указанные в презентации 15,5 млн семейных ячеек с детьми до 18 лет [2] служат общерыночным ориентиром и в расчёте не используются.")
para(doc, "Чек 5 500 ₽ пока не подкреплён заказами продукта. Для сравнения: СберАналитика оценила средний чек на детские товары в 1 522 ₽ [3], а по данным эквайринга Банка Русский Стандарт средний чек маркетплейсов составил 3 008 ₽, онлайн-магазинов детских товаров — 1 843 ₽ [4]. Эти покупки отличаются от заранее выбранного подарка, поэтому показатели годятся для стресс-теста, но не определяют чек нашего сценария.")
para(doc, "Доли 15%, 45% и 60% не имеют внешнего подтверждения. Первые две нужно измерить по воронке подключения и покупки. Долю действительно дополнительного оборота — определить с контрольной группой, учитывая заказы, которые семья совершила бы и без продукта. Маржу 8% и бюджет 18 + 12 млн ₽ должны подтвердить финансовые и продуктовые команды с учётом возвратов, скидок, логистики, поддержки, безопасности детских данных и интеграций.")

doc.add_paragraph("Данные для пилота", "Heading 1")
for text in [
    "Уникальные семьи за активными детскими аккаунтами, доступность взрослого для подключения и отдельная база подростков.",
    "Воронка: показ предложения → привязка → первая одобренная цель → достижение порога → запрос подарка → оплаченный заказ → возврат.",
    "Распределение цен целей и фактический чек оплаченных подарков по возрастам и категориям.",
    "Разница в чистом обороте и маржинальном вкладе между тестовой и контрольной группами.",
    "Смета разработки и ежемесячные расходы эксплуатации, включая интеграции, поддержку и комплаенс.",
]:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(4)

doc.add_paragraph("Публичные источники", "Heading 1")
for source in [
    "[1] Сбер, результаты за пять месяцев 2026 года: https://e-disclosure.azipi.ru/messages/4551199/",
    "[2] Семья и дети в России, перепись 2020 года: https://files.oprf.ru/storage/documents/doklad-semiya-deti-2024.pdf",
    "[3] СберАналитика, детские товары в 2025 году: https://sberanalytics.ru/researches/children-goods-spending",
    "[4] Банк Русский Стандарт, покупки онлайн в 2025 году: https://www.rsb.ru/press-center/publications/2026/180126/",
]:
    p = doc.add_paragraph(source)
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        run.font.size = Pt(8.5)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Семейный конструктор целей · расчёт экономической модели · 30.09.2026")
for run in footer.runs:
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(106, 117, 114)

doc.save(OUTPUT)
print(OUTPUT)
